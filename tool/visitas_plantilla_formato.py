"""Genera assets/visitas_plantilla_formato.xlsx (26 sep 2026).

La plantilla que se descarga en Visitas > Formatos. Los encabezados y los
nombres de los tipos DEBEN coincidir con lib/visitas/visitas_formato_excel.dart
(kColumnasPreguntas, kColumnasTablas) y lib/visitas/visitas_models.dart
(kItemTiposLabel, kEscalasLabel): la prueba
test/visitas/visitas_formato_excel_test.dart lo verifica.

Uso:  pip install openpyxl && python3 tool/visitas_plantilla_formato.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = Path(__file__).resolve().parent.parent / "assets" / "visitas_plantilla_formato.xlsx"

MORADO = "7C3AED"
MORADO_CLARO = "EDE9FE"
AMARILLO = "FEF3C7"

COLUMNAS_PREGUNTAS = [
    ("Sección", 24),
    ("Pregunta", 60),
    ("Tipo de respuesta", 34),
    ("Opciones (separadas por ;)", 30),
    ("Obligatoria", 13),
    ("Foto obligatoria si no cumple", 17),
    ("Cantidad esperada", 18),
    ("Ayuda para el profesional", 40),
]

COLUMNAS_TABLAS = [
    ("Tabla", 26),
    ("Cada fila es un(a)", 18),
    ("Columnas para escribir (separadas por ;)", 42),
    ("Columnas para calificar (separadas por ;)", 46),
    ("Escala", 34),
    ("Filas fijas (separadas por ;)", 40),
]

TIPOS = [
    ("Cumple / No cumple / No aplica",
     "Se califica 1 (cumple), 0 (no cumple) o NA. Cuenta en el % de cumplimiento y cada \"No cumple\" se vuelve tarea."),
    ("Sí / No",
     "Solo Sí o No, sin \"No aplica\". Un \"No\" cuenta como no cumple y se vuelve tarea."),
    ("Sí / No con cantidad y vencimiento",
     "Para revisar que un elemento esté (gasas del botiquín, guantes…) y anotar cuántos hay y cuándo vencen. Escribe lo esperado en \"Cantidad esperada\" (ej. \"1 paquete x 20\")."),
    ("Lista de opciones",
     "El profesional elige una de las opciones que escribas en \"Opciones\", separadas por punto y coma. No califica."),
    ("Respuesta corta", "Un dato corto: un nombre, una placa. No califica."),
    ("Párrafo", "Texto largo o una descripción. No califica."),
    ("Número", "Una cantidad o medida: comensales, temperatura. No califica."),
    ("Fecha", "Una fecha: último fumigado, vencimiento del certificado. No califica."),
]

ESCALAS = [
    "Bueno / Malo / Regular / No cuenta",
    "Cumple / No cumple / No aplica",
    "Sí / No",
]

EJEMPLO_PREGUNTAS = [
    ["Recepción", "Nombre de quien recibe la mercancía", "Respuesta corta", "", "Sí", "", "", ""],
    ["Recepción", "Número de comensales del día", "Número", "", "Sí", "", "", "Según la planilla del establecimiento"],
    ["Cocina", "¿Los pisos, paredes y techos están limpios?", "Cumple / No cumple / No aplica", "", "", "Sí", "", "Mira debajo de mesones y estufas"],
    ["Cocina", "¿La cadena de frío está documentada?", "Cumple / No cumple / No aplica", "", "", "No", "", ""],
    ["Cocina", "¿Hay control de plagas vigente?", "Sí / No", "", "", "Sí", "", ""],
    ["Cocina", "Fecha del último fumigado", "Fecha", "", "No", "", "", ""],
    ["Cocina", "Estado general de la cocina", "Lista de opciones", "Excelente; Bueno; Regular; Malo", "Sí", "", "", ""],
    ["Botiquín", "Gasas estériles", "Sí / No con cantidad y vencimiento", "", "", "No", "1 paquete x 20", ""],
    ["Botiquín", "Guantes de látex", "Sí / No con cantidad y vencimiento", "", "", "No", "10 pares", ""],
    ["Cierre", "Descripción de lo observado", "Párrafo", "", "No", "", "", "Lo que no quedó en las preguntas"],
]

EJEMPLO_TABLAS = [
    ["Extintores", "Extintor", "Ubicación; Tipo; Capacidad; Fecha próxima carga",
     "Señalización; Estado del cilindro; Manómetro; Pasador; Manguera", ESCALAS[0], ""],
    ["Áreas del establecimiento", "Área", "",
     "Limpieza; Orden; Iluminación", ESCALAS[1], "Cocina; Bodega; Comedor; Baños"],
]

BORDE = Border(*(Side(style="thin", color="C4B5FD"),) * 4)


def encabezados(ws, columnas, fila=1):
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MORADO)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDE
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[fila].height = 34
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def lista(ws, rango, valores):
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(valores) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Valor no válido",
        error="Elige un valor de la lista.",
    )
    ws.add_data_validation(dv)
    dv.add(rango)


def hoja_preguntas(ws, filas=None):
    encabezados(ws, COLUMNAS_PREGUNTAS)
    for r, fila in enumerate(filas or [], start=2):
        for c, v in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=v or None)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.border = BORDE
    lista(ws, "C2:C500", [t for t, _ in TIPOS])
    lista(ws, "E2:E500", ["Sí", "No"])
    lista(ws, "F2:F500", ["Sí", "No"])


def hoja_tablas(ws, filas=None):
    encabezados(ws, COLUMNAS_TABLAS)
    for r, fila in enumerate(filas or [], start=2):
        for c, v in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=v or None)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.border = BORDE
    lista(ws, "E2:E100", ESCALAS)


def instrucciones(ws):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    ws.sheet_properties.tabColor = MORADO
    fila = 1

    def titulo(texto):
        nonlocal fila
        c = ws.cell(row=fila, column=1, value=texto)
        c.font = Font(bold=True, size=13, color=MORADO)
        fila += 1

    def par(k, v, fondo=None):
        nonlocal fila
        a = ws.cell(row=fila, column=1, value=k)
        b = ws.cell(row=fila, column=2, value=v)
        a.font = Font(bold=True)
        for c in (a, b):
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDE
            if fondo:
                c.fill = PatternFill("solid", fgColor=fondo)
        fila += 1

    titulo("Plantilla para formatos de visita")
    par("Cómo se usa",
        "Escribe una pregunta por fila en la hoja \"Preguntas\" y, si el formato lleva tablas (extintores, "
        "áreas del establecimiento), una tabla por fila en la hoja \"Tablas\". No cambies los encabezados. "
        "Luego en la app: Visitas > Formatos > Crear desde Excel. Queda como borrador para que lo revises.",
        AMARILLO)
    par("¿Ya tienes el formato en otro Excel?",
        "No hace falta pasarlo a esta plantilla: súbelo tal cual en \"Crear desde Excel\". La app busca la columna "
        "de las preguntas y las secciones, y te lo muestra en el editor para que corrijas lo que haga falta antes de guardar.")
    par("Ejemplos",
        "Las hojas \"Ejemplo preguntas\" y \"Ejemplo tablas\" muestran un formato lleno. Son solo de muestra: "
        "la app no las importa.")
    par("Área y nombre", "Se eligen en la aplicación al importar el archivo.")
    fila += 1
    titulo("Hoja Preguntas")
    par("Sección", "Opcional. Agrupa las preguntas que van juntas (Cocina, Bodega, Botiquín). En la visita, cada sección es una página.")
    par("Pregunta", "Obligatoria. Lo que se revisa, como se le pregunta al profesional.")
    par("Tipo de respuesta", "Elige de la lista. Si lo dejas vacío queda \"Cumple / No cumple / No aplica\".")
    par("Opciones (separadas por ;)", "Solo para \"Lista de opciones\": Excelente; Bueno; Regular; Malo.")
    par("Obligatoria", "Sí o No. Las que califican siempre son obligatorias; las demás pueden quedar opcionales.")
    par("Foto obligatoria si no cumple", "Sí exige foto cuando la respuesta es \"No cumple\" o \"No\". No la deja opcional.")
    par("Cantidad esperada", "Solo para \"Sí / No con cantidad y vencimiento\": 1 unidad, paquete x 20.")
    par("Ayuda para el profesional", "Opcional. Sale debajo de la pregunta: qué mirar o cómo medir.")
    fila += 1
    titulo("Tipos de respuesta")
    for t, ayuda in TIPOS:
        par(t, ayuda, MORADO_CLARO)
    fila += 1
    titulo("Hoja Tablas")
    par("Tabla", "Nombre de la tabla (Extintores).")
    par("Cada fila es un(a)", "Cómo se llama cada fila (Extintor, Área).")
    par("Columnas para escribir", "Datos que se escriben en cada fila, separados por punto y coma: Ubicación; Tipo; Capacidad.")
    par("Columnas para calificar", "Lo que se califica en cada fila, separado por punto y coma: Señalización; Manómetro; Pasador.")
    par("Escala", "Con qué se califica: " + " · ".join(ESCALAS) + ". Vacío = la de extintores (B/M/R/NC).")
    par("Filas fijas", "Opcional. Si las escribes (Cocina; Bodega; Baños) el profesional califica esas filas, como una "
        "cuadrícula de Google Forms. Si la dejas vacía, él agrega una fila por cada equipo que encuentre.")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    instrucciones(ws)
    hoja_preguntas(wb.create_sheet("Preguntas"))
    hoja_tablas(wb.create_sheet("Tablas"))
    ej = wb.create_sheet("Ejemplo preguntas")
    ej.sheet_properties.tabColor = "F59E0B"
    hoja_preguntas(ej, EJEMPLO_PREGUNTAS)
    ejt = wb.create_sheet("Ejemplo tablas")
    ejt.sheet_properties.tabColor = "F59E0B"
    hoja_tablas(ejt, EJEMPLO_TABLAS)
    wb.active = 0
    wb.save(SALIDA)
    print(f"Plantilla escrita en {SALIDA}")


if __name__ == "__main__":
    main()
